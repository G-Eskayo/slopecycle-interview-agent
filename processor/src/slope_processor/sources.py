"""Turn a Source (Drive file) into numbered lines the model can cite and the validator can check."""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass
from enum import Enum
from pathlib import PurePath


class SourceKind(Enum):
    TEXT = "text"
    DOCX = "docx"
    AUDIO = "audio"
    FOLDER = "folder"
    SUBMISSION = "submission"
    UNSUPPORTED = "unsupported"


class UnsupportedSource(Exception):
    """The file can't be read as an interview. The message is shown to teammates in Run Log."""


TEXT_MIMES = {"application/vnd.google-apps.document", "text/plain", "text/markdown", "text/x-markdown"}
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
AUDIO_EXTS = {".m4a", ".mp3", ".wav", ".aac", ".caf", ".mp4", ".mov", ".aiff", ".aif", ".flac", ".ogg", ".webm"}
TEXT_EXTS = {".txt", ".md", ".markdown"}


SUBMISSION_MIME = "application/x-slopecycle-submission"


def classify(name: str, mime: str) -> SourceKind:
    ext = PurePath(name).suffix.lower()
    if mime == SUBMISSION_MIME:
        return SourceKind.SUBMISSION
    if mime == "application/vnd.google-apps.folder":
        return SourceKind.FOLDER
    if mime in TEXT_MIMES or ext in TEXT_EXTS:
        return SourceKind.TEXT
    if mime == DOCX_MIME or ext == ".docx":
        return SourceKind.DOCX
    if mime.startswith("audio/") or ext in AUDIO_EXTS:
        return SourceKind.AUDIO
    return SourceKind.UNSUPPORTED


@dataclass(frozen=True)
class Line:
    number: int
    text: str
    timestamp: str | None = None
    section: str | None = None  # "notes" / "recording" when one Submission holds both

    @property
    def locator(self) -> str:
        where = f"{self.section}, line {self.number}" if self.section else f"line {self.number}"
        return where + (f" ({self.timestamp})" if self.timestamp else "")


def text_from_bytes(data: bytes, kind: SourceKind) -> str:
    if kind is SourceKind.DOCX:
        return _docx_text(data)
    if kind is not SourceKind.TEXT:
        raise UnsupportedSource(f"{kind.value} files are not read as text")
    if b"\x00" in data[:4096] or _control_ratio(data[:4096]) > 0.1:
        raise UnsupportedSource("file is not text (looks binary or corrupt)")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _control_ratio(sample: bytes) -> float:
    if not sample:
        return 0.0
    bad = sum(1 for b in sample if b < 32 and b not in (9, 10, 13))
    return bad / len(sample)


def _docx_text(data: bytes) -> str:
    import docx  # imported here so text-only installs don't pay for it

    try:
        document = docx.Document(io.BytesIO(data))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise UnsupportedSource(f"Word file could not be read ({type(exc).__name__})") from exc
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def lines_from_text(text: str, section: str | None = None, start: int = 1) -> list[Line]:
    out: list[Line] = []
    for raw in text.splitlines():
        stripped = raw.strip()
        if stripped:
            out.append(Line(start + len(out), stripped, None, section))
    return out


def _clock(seconds: float) -> str:
    s = int(seconds)
    return f"{s // 3600:02d}:{s % 3600 // 60:02d}:{s % 60:02d}"


def lines_from_transcript(transcript: dict, section: str | None = None, start: int = 1) -> list[Line]:
    out: list[Line] = []
    for seg in transcript.get("segments", []):
        text = " ".join(str(seg.get("text", "")).split())
        if not text:
            continue
        ts = _clock(float(seg.get("start", 0)))
        out.append(Line(start + len(out), f"[{ts}] {seg.get('speaker', 'S1')}: {text}", ts, section))
    return out
