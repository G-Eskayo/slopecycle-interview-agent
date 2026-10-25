import io

import docx
import pytest

from slope_processor.sources import (
    SourceKind,
    UnsupportedSource,
    classify,
    lines_from_text,
    lines_from_transcript,
    text_from_bytes,
)


@pytest.mark.parametrize(
    "name,mime,kind",
    [
        ("notes", "application/vnd.google-apps.document", SourceKind.TEXT),
        ("notes.txt", "text/plain", SourceKind.TEXT),
        ("notes.md", "text/markdown", SourceKind.TEXT),
        ("notes.docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document", SourceKind.DOCX),
        ("call.m4a", "audio/x-m4a", SourceKind.AUDIO),
        ("call.mp3", "audio/mpeg", SourceKind.AUDIO),
        ("call.MOV", "video/quicktime", SourceKind.AUDIO),
        ("weird", "audio/ogg", SourceKind.AUDIO),
        ("scan.pdf", "application/pdf", SourceKind.UNSUPPORTED),
        ("photo.jpg", "image/jpeg", SourceKind.UNSUPPORTED),
        ("sheet", "application/vnd.google-apps.spreadsheet", SourceKind.UNSUPPORTED),
        ("sub", "application/vnd.google-apps.folder", SourceKind.FOLDER),
        ("no-extension", "application/octet-stream", SourceKind.UNSUPPORTED),
    ],
)
def test_classify(name, mime, kind):
    assert classify(name, mime) is kind


def test_text_lines_are_numbered_and_skip_blanks():
    lines = lines_from_text("First line\n\n   \nSecond line\r\nThird")
    assert [(l.number, l.text) for l in lines] == [(1, "First line"), (2, "Second line"), (3, "Third")]
    assert lines[1].locator == "line 2"


def test_utf8_bom_and_latin1_fallback():
    assert text_from_bytes("﻿hi".encode("utf-8"), SourceKind.TEXT) == "hi"
    assert text_from_bytes("café".encode("latin-1"), SourceKind.TEXT) == "café"


def test_docx_paragraphs_and_tables():
    d = docx.Document()
    d.add_paragraph("Parent said kids outgrow jackets yearly.")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Q: spend?"
    t.cell(0, 1).text = "About $300"
    buf = io.BytesIO()
    d.save(buf)
    text = text_from_bytes(buf.getvalue(), SourceKind.DOCX)
    assert "outgrow jackets" in text and "About $300" in text


def test_corrupt_docx_is_unsupported_not_a_crash():
    with pytest.raises(UnsupportedSource, match="could not be read"):
        text_from_bytes(b"PK\x03\x04 not really a zip", SourceKind.DOCX)


def test_binary_garbage_in_a_txt_is_rejected():
    with pytest.raises(UnsupportedSource, match="not text"):
        text_from_bytes(bytes(range(256)) * 4, SourceKind.TEXT)


def test_transcript_lines_carry_speaker_and_timestamp():
    lines = lines_from_transcript(
        {"segments": [
            {"speaker": "S1", "start": 3.2, "end": 7.0, "text": "How often do they outgrow jackets?"},
            {"speaker": "S2", "start": 3725.9, "end": 3730.0, "text": "  Every single year.  "},
            {"speaker": "S2", "start": 3731.0, "end": 3731.5, "text": "   "},
        ]}
    )
    assert [l.text for l in lines] == ["[00:00:03] S1: How often do they outgrow jackets?", "[01:02:05] S2: Every single year."]
    assert lines[1].locator == "line 2 (01:02:05)"
